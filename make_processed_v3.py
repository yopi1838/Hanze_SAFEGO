"""make_processed_v3.py -- rebuild the test base shear with the group's formula (Ihsan / Eleni, 26 Oct 2026).

    V = g/1000 * (165 a13 + 165 a14) + 12.94 kN * (f1 a15 + f2 a16 + f3 a17)            [kN, accelerations in g]
        a13 / a14 : slab accelerometers (0.49 / 2.28 m), 165 kg per slab
        a15-a17   : wall accelerometers at h1 / h2 / h3 = 0.60 / 1.26 / 2.06 m (read from TestX_Info.xlsx, F16-F18)
        f1 = (h1 + (h2-h1)/2)/H,  f2 = ((h2-h1)/2 + (h3-h2)/2)/H,  f3 = (H - h3 + (h3-h2)/2)/H,  H = 2.58 m
           = 0.360 / 0.283 / 0.357  ->  475.5 / 373.2 / 470.3 kg of the 1319 kg wall (12.94 kN, g = 9.8104)
        top steel beam (61.4 kg on a17): excluded, as in the group's script (--beam to include it)
    All channels ABSOLUTE as recorded (ch12, the table, is NOT added: that was the error of the old processed files,
    which inflated the peaks by about 35 %).

Written in the layout every reader here expects (postprocess_r2tm_full25.load_test_run, exp_envelope_peaks.py):
    Time [s] | U_avg [m] = 0.5 (ch3 + ch4) - Z0, Z0 at the first sample of RUN 1 (global zero, NOT sign-flipped)
    | base_shear_N | base_shear_kN = +V (the readers apply sign -1: restoring force positive with positive U_avg)
    | mass_scale = 1 | m13 ... m17 (kg)
The group's plot flips the displacement instead of the force and zeroes it per run; the loops are the same.
Offsets, the US-1 run-18 channel slip and the 25 Hz low-pass are left to the readers (they do all three).

EXP_DATA is only read. Output goes to a folder of its own.

USAGE   python make_processed_v3.py <raw_dir> <out_dir> [--beam]
        python make_processed_v3.py EXP_DATA exp_processed_v3
"""
import sys, os, re, glob
import numpy as np

G = 9.8104
H = 2.58
WALL_KN = 12.94
SLAB_KG = (165.0, 165.0)
BEAM_KG = 61.4
TESTS = ("Test9", "Test12")


def heights(raw_dir, test):
    import openpyxl
    p = os.path.join(raw_dir, "{}_Info.xlsx".format(test))
    try:
        ws = openpyxl.load_workbook(p, read_only=True, data_only=True).worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        h = tuple(float(rows[r][5]) for r in (15, 16, 17))
    except Exception as e:
        print("  {}: heights not read ({}); using 0.60 / 1.26 / 2.06 m".format(p, e))
        h = (0.60, 1.26, 2.06)
    if not (0 < h[0] < h[1] < h[2] < H):
        raise ValueError("{}: invalid accelerometer heights {}".format(test, h))
    return h


def masses(h, beam):
    h1, h2, h3 = h
    f = ((h1 + (h2 - h1) / 2) / H, ((h2 - h1) / 2 + (h3 - h2) / 2) / H, (H - h3 + (h3 - h2) / 2) / H)
    mw = WALL_KN * 1e3 / G
    m = {13: SLAB_KG[0], 14: SLAB_KG[1], 15: mw * f[0], 16: mw * f[1], 17: mw * f[2] + (BEAM_KG if beam else 0.0)}
    return m, f


def load_raw(path):
    import openpyxl
    ws = openpyxl.load_workbook(path, read_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True); hdr = next(it)
    col = {}
    for j, hh in enumerate(hdr):
        if hh is None:
            continue
        if isinstance(hh, str) and hh.strip().lower() == "time":
            col["time"] = j
        else:
            try:
                col[int(float(hh))] = j
            except (TypeError, ValueError):
                pass
    need = ("time", 3, 4, 13, 14, 15, 16, 17)
    miss = [k for k in need if k not in col]
    if miss:
        raise KeyError("{}: columns {} not found in row 1".format(os.path.basename(path), miss))
    rows = [r for r in it if r[col["time"]] is not None]
    A = {k: np.array([r[j] if (j < len(r) and r[j] is not None) else np.nan for r in rows], float) for k, j in col.items()}
    return A


def run_no(path):
    m = re.search(r"Run0*(\d+)", os.path.basename(path), re.I)
    return int(m.group(1)) if m else None


def main():
    if len(sys.argv) < 3:
        print(__doc__); sys.exit(1)
    raw_dir, out_dir = sys.argv[1], sys.argv[2]
    beam = "--beam" in sys.argv
    if os.path.abspath(out_dir).startswith(os.path.abspath(raw_dir) + os.sep) or os.path.abspath(out_dir) == os.path.abspath(raw_dir):
        raise SystemExit("out_dir must not be inside {} (raw data is read-only)".format(raw_dir))
    os.makedirs(out_dir, exist_ok=True)
    import openpyxl
    for test in TESTS:
        files = sorted([f for f in glob.glob(os.path.join(raw_dir, test + "Run*.xlsx"))
                        if "processed" not in f.lower() and run_no(f)], key=run_no)
        if not files:
            print("{}: no raw files in {}".format(test, raw_dir)); continue
        h = heights(raw_dir, test)
        m, f = masses(h, beam)
        A1 = load_raw([p for p in files if run_no(p) == 1][0])
        u1 = 0.5 * (A1[3] + A1[4]); Z0 = float(u1[np.isfinite(u1)][0])
        print("{}: heights {} m, fractions {:.3f}/{:.3f}/{:.3f}, masses ".format(test, h, *f)
              + ", ".join("m{} {:.1f}".format(k, v) for k, v in sorted(m.items()))
              + " kg (sum {:.1f}{}), Z0 {:.6f} m".format(sum(m.values()), ", beam on a17" if beam else ", no beam", Z0))
        for p in files:
            n = run_no(p)
            A = load_raw(p)
            ok = np.isfinite(A["time"]) & np.isfinite(A[3]) & np.isfinite(A[4]) & np.all([np.isfinite(A[k]) for k in m], axis=0)
            t = A["time"][ok]
            U = 0.5 * (A[3][ok] + A[4][ok]) - Z0
            V = G * sum(m[k] * A[k][ok] for k in m)                      # N
            wb = openpyxl.Workbook(write_only=True); ws = wb.create_sheet("processed")
            ws.append(["Time", "U_avg", "base_shear_N", "base_shear_kN", "mass_scale", "m13_kg", "m14_kg",
                       "m15_kg", "m16_kg", "m17_kg", "formula"])
            meta = [1.0] + [round(m[k], 3) for k in (13, 14, 15, 16, 17)] + [
                "group 26 Oct 2026: slabs on a13/a14, wall 12.94 kN by sensor height; absolute; {}".format(
                    "beam 61.4 kg on a17" if beam else "no beam")]
            for i in range(len(t)):
                row = [float(t[i]), float(U[i]), float(V[i]), float(V[i] / 1e3)]
                ws.append(row + meta if i == 0 else row + [1.0])
            outp = os.path.join(out_dir, "{}Run{}_processed_globalzero.xlsx".format(test, n))
            wb.save(outp)
            print("  run {:2d}: {} samples, |V|max {:.2f} kN (raw)".format(n, len(t), float(np.max(np.abs(V))) / 1e3))
    print("->", out_dir)


if __name__ == "__main__":
    main()
