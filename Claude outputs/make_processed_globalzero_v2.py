"""make_processed_globalzero_v2.py -- rebuild TestXXRunYY_processed_globalzero.xlsx from the raw SAFEGo channel files.

v2 (5 Oct 2026) replaces the recipe that produced the earlier processed files:
    OLD : m = midpoint tributary split of a 1133 kg wall (rho 1619) x 2.317348, the factor having been
          CALIBRATED so that Run 24 of US-1 reaches the 29.5 kN reported in the paper; and the wall
          channels were taken as RELATIVE to the table, with ch12 added to make them absolute.
    v2  : masses from the group's sheet -- wall 13.16 kN at gamma 18.8 kN/m3 (1340 kg) split 1/4-1/2-1/4
          between the top-quarter / middle / bottom-quarter accelerometers, + 120 kg slab at each quarter
          level, + 151.4 kg loading beam lumped into the top quarter; springs (10.35 kN) excluded.
          No factor. The wall channels 15/16/17 are ABSOLUTE as recorded (check_accel_abs_vs_rel.py,
          Test9Run24 and Test12Run24: all three track ch12 in the 0.5-3 Hz band), so ch12 is NOT added.

    U_avg [m]      = 0.5 (ch3 + ch4) - Z0        Z0 = 0.5 (ch3 + ch4) at the first finite sample of RUN 1
                                                 (ch5 is NOT subtracted: the potentiometer frames ride on the table)
    base_shear_N   = g0 * (m15 ch15 + m16 ch16 + m17 ch17)       channels in g, absolute

Raw file layout: row 1 = [None, 'Time', 1, 2, ..., 17]; data from row 2.
Usage:
    python make_processed_globalzero_v2.py Test9  <raw_dir> <out_dir>
    python make_processed_globalzero_v2.py Test12 <raw_dir> <out_dir>
    python make_processed_globalzero_v2.py Test9  <raw_dir> <out_dir> --compare <old_processed_dir>
        (prints, per run, old vs new peak base shear so the change is traceable)
    python make_processed_globalzero_v2.py Test9  <raw_dir> <out_dir_slabs> --slabs
        five-term formula: the slab and beam masses ride on THEIR OWN accelerometers instead of the wall's
            base_shear_N = g0 * (335.03 ch15 + 670.07 ch16 + 335.03 ch17 + 120 ch13 + 271.44 ch14)
        ch13 = slab 1 (0.49 m), ch14 = slab 2 (2.28 m; the loading beam sits above it, so its 151.4 kg go there too).
        Same total mass (1731.6 kg); differs from the three-term version only when the joists move
        relative to the wall. Use a separate out_dir and compare_processed_old_new.py to see the difference.
Writes <out_dir>/TestXXRunYY_processed_globalzero.xlsx with the masses and the convention recorded in every file.
"""
import sys, os, re, glob
import numpy as np
import openpyxl

G0 = 9.81
MASS_SCALE = 1.0
# three-term (group's sheet): slab and beam masses lumped onto the wall quarter sensors
M3 = {15: 455.03, 16: 670.07, 17: 606.47}
# five-term: wall-only masses on 15/16/17, slab 1 on ch13, slab 2 + beam on ch14
M5 = {15: 335.03, 16: 670.07, 17: 335.03, 13: 120.0, 14: 271.44}
USE_SLABS = "--slabs" in sys.argv
M = M5 if USE_SLABS else M3
ACCEL_CONVENTION = ("absolute as recorded; slabs on ch13/ch14, beam on ch14" if USE_SLABS
                    else "absolute as recorded (ch12 not added); slabs+beam lumped on wall sensors")

def load_raw(path):
    ws = openpyxl.load_workbook(path, read_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True); hdr = next(it)
    col = {}
    for j, h in enumerate(hdr):
        if h is None: continue
        if isinstance(h, str) and h.strip().lower() == "time": col["time"] = j
        else:
            try: col[int(float(h))] = j
            except (TypeError, ValueError): pass
    for k in ("time", 3, 4, 15, 16, 17) + ((13, 14) if USE_SLABS else ()):
        if k not in col: raise KeyError("{}: column {} not found in row 1".format(os.path.basename(path), k))
    rows = [r for r in it if r[col["time"]] is not None]
    A = np.array([[r[j] if (j < len(r) and r[j] is not None) else np.nan for j in range(len(hdr))] for r in rows], float)
    return A[:, col["time"]], {k: A[:, j] for k, j in col.items() if k != "time"}

def run_no(path):
    m = re.search(r"Run0*(\d+)", os.path.basename(path), re.I); return int(m.group(1)) if m else None

def main():
    if len(sys.argv) < 4:
        print(__doc__); sys.exit(1)
    test, raw_dir, out_dir = sys.argv[1], sys.argv[2], sys.argv[3]
    cmp_dir = sys.argv[sys.argv.index("--compare") + 1] if "--compare" in sys.argv else None
    sys.argv = [a for a in sys.argv if a != "--slabs"]
    os.makedirs(out_dir, exist_ok=True)
    files = sorted([f for f in glob.glob(os.path.join(raw_dir, test + "Run*.xlsx")) if "processed" not in f.lower() and run_no(f)], key=run_no)
    if not files: raise FileNotFoundError("no {}RunNN.xlsx in {}".format(test, raw_dir))
    # global zero from run 1
    r1 = [f for f in files if run_no(f) == 1]
    if not r1: raise FileNotFoundError("{}Run1.xlsx is needed for the global zero".format(test))
    t1, ch1 = load_raw(r1[0]); u1 = 0.5 * (ch1[3] + ch1[4]); Z0 = float(u1[np.isfinite(u1)][0])
    print("{}: global zero Z0 = 0.5(ch3+ch4) at t0 of run 1 = {:.8f} m".format(test, Z0))
    print("masses " + ", ".join("ch{}: {:.2f}".format(k, M[k]) for k in sorted(M)) + " kg (sum {:.1f}), scale {}, {}".format(sum(M.values()), MASS_SCALE, ACCEL_CONVENTION))
    for f in files:
        n = run_no(f); t, ch = load_raw(f)
        U = 0.5 * (ch[3] + ch[4]) - Z0
        V = G0 * MASS_SCALE * sum(M[k] * ch[k] for k in M)
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(["Time", "U_avg", "base_shear_N", "base_shear_kN", "mass_scale", "m15_kg", "m16_kg", "m17_kg", "m13_kg", "m14_kg", "accel_convention", "Z0_m"])
        for i in range(len(t)):
            ws.append([float(t[i]), float(U[i]), float(V[i]), float(V[i]) / 1e3, MASS_SCALE, M[15], M[16], M[17], M.get(13, 0.0), M.get(14, 0.0),
                       ACCEL_CONVENTION if i == 0 else None, Z0 if i == 0 else None])
        out = os.path.join(out_dir, "{}Run{}_processed_globalzero.xlsx".format(test, n)); wb.save(out)
        line = "  run {:2d}: peak U {:7.3f} mm  peak V {:6.2f} kN".format(n, np.nanmax(np.abs(U)) * 1e3, np.nanmax(np.abs(V)) / 1e3)
        if cmp_dir:
            old = os.path.join(cmp_dir, "{}Run{}_processed_globalzero.xlsx".format(test, n))
            if os.path.isfile(old):
                wso = openpyxl.load_workbook(old, read_only=True).worksheets[0]
                vo = np.array([r[3] for r in wso.iter_rows(min_row=2, values_only=True) if r[3] is not None], float)
                line += "   (old file: {:6.2f} kN, old/new {:.2f})".format(np.nanmax(np.abs(vo)), np.nanmax(np.abs(vo)) / (np.nanmax(np.abs(V)) / 1e3))
        print(line)
    print("-> {} files in {}".format(len(files), out_dir))

if __name__ == "__main__":
    main()
