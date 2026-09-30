# -*- coding: ascii -*-
"""
make_processed_globalzero.py -- rebuild TestXXRunNN_processed_globalzero.xlsx
from the raw TestXXRunNN.xlsx, exactly as Test9Run1_processed_globalzero.xlsx
was made. Verified against that file to machine precision (see VERIFY).

WHAT THE PROCESSED FILE IS (recovered by least squares on Test9Run1, residual
~1e-15 m and ~1e-13 N -- these are the formulas, not approximations):

    U_avg [m]      = 0.5 (ch3 + ch4) - Z0
                     ch3, ch4 = top-quarter left / right potentiometers, already
                     relative to the table (the frames ride on it), so ch5 is
                     NOT subtracted. Z0 = 0.5 (ch3 + ch4) at t = 0 of RUN 1:
                     "globalzero" -- every run of a test uses run 1's zero, so a
                     later run's U_avg carries the residual of all runs before it.
    base_shear_N   = g [ (m15 + m16 + m17) ch12 + m15 ch15 + m16 ch16 + m17 ch17 ]
                     ch12 = table acceleration [g], ch15/16/17 = wall accelerometers
                     at the bottom quarter / mid / top quarter [g], read as RELATIVE
                     to the table; absolute = ch12 + chi. So this is sum(m_i a_i,abs).
    base_shear_kN  = base_shear_N / 1000
    mass_scale, m15_kg, m16_kg, m17_kg : constants, copied.
                     m15 946.679, m16 743.092, m17 936.500 kg (sum 2626.27);
                     mass_scale 2.317348 = sum / 1133.31 kg. The unscaled tributary
                     masses (408.5 / 320.7 / 404.1) are not derivable from the
                     Info sheet (its 'Associated Mass' column is empty), so the
                     constants are carried as found.

TEST 12 (US-2) is the twin wall: same geometry, same joists, same overburden,
same channel table (Test12_Info.xlsx differs from Test9_Info.xlsx only in the
bottom-quarter gauge height, 0.60 vs 0.66 m, which enters nothing here). The
same masses and mass_scale are therefore applied. That is an ASSUMPTION about
how the original processing was done for Test 12 -- flagged in the output.

Usage:
    python make_processed_globalzero.py Test12 [RAW_DIR] [OUT_DIR]
        processes every Test12RunNN.xlsx in RAW_DIR (default .), run 1 first
        so its zero is known; writes Test12RunNN_processed_globalzero.xlsx
    python make_processed_globalzero.py --verify
        rebuilds Test9Run1 and diffs it against the supplied processed file
"""
import os, re, sys, glob
import numpy as np
import openpyxl

M15, M16, M17 = 946.6793249800786, 743.0923733714598, 936.4999773996476
MASS_SCALE    = 2.317348156486243
G             = 9.81
COLS = ["Time", "U_avg", "base_shear_N", "base_shear_kN", "mass_scale", "m15_kg", "m16_kg", "m17_kg"]
RUN_RE = re.compile(r"^(Test\d+)Run(\d+)\.xlsx$")

def load_raw(path):
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True); hdr = next(rows)
    d = np.array([r for r in rows if r and r[1] is not None], dtype=float)
    # layout: (index, Time, ch1 .. ch17)
    t = d[:, 1]; ch = {i: d[:, 1 + i] for i in range(1, 18)}
    return t, ch

def process(t, ch, z0):
    u = 0.5 * (ch[3] + ch[4]) - z0
    F = G * ((M15 + M16 + M17) * ch[12] + M15 * ch[15] + M16 * ch[16] + M17 * ch[17])
    n = len(t)
    return np.c_[t, u, F, F / 1000.0, np.full(n, MASS_SCALE), np.full(n, M15), np.full(n, M16), np.full(n, M17)]

def write_xlsx(path, arr):
    wb = openpyxl.Workbook(write_only=True); ws = wb.create_sheet("Sheet1")
    ws.append(COLS)
    for row in arr:
        ws.append([float(x) for x in row])
    wb.save(path)

def zero_of_run1(raw_dir, test):
    p = os.path.join(raw_dir, "{}Run1.xlsx".format(test))
    if not os.path.isfile(p):
        raise SystemExit("{} not found -- run 1 defines the global zero and must be present".format(p))
    t, ch = load_raw(p)
    return float(0.5 * (ch[3][0] + ch[4][0]))

def verify(raw_dir="."):
    t, ch = load_raw(os.path.join(raw_dir, "Test9Run1.xlsx"))
    z0 = float(0.5 * (ch[3][0] + ch[4][0]))
    mine = process(t, ch, z0)
    ws = openpyxl.load_workbook(os.path.join(raw_dir, "Test9Run1_processed_globalzero.xlsx"), read_only=True, data_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True); hdr = list(next(rows))
    ref = np.array([r for r in rows if r and r[0] is not None], dtype=float)
    print("header match :", hdr == COLS)
    print("shape        : mine {} ref {}".format(mine.shape, ref.shape))
    for j, c in enumerate(COLS):
        err = np.max(np.abs(mine[:, j] - ref[:, j])); scale = max(np.max(np.abs(ref[:, j])), 1e-30)
        print("  {:14s} max |diff| {:.3e}   (rel {:.1e})".format(c, err, err / scale))
    print("global zero Z0 (Test9) = {:.8f} m".format(z0))

def main():
    if len(sys.argv) >= 2 and sys.argv[1] == "--verify":
        verify(sys.argv[2] if len(sys.argv) > 2 else "."); return
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    test = sys.argv[1]; raw_dir = sys.argv[2] if len(sys.argv) > 2 else "."; out_dir = sys.argv[3] if len(sys.argv) > 3 else raw_dir
    os.makedirs(out_dir, exist_ok=True)
    z0 = zero_of_run1(raw_dir, test)
    print("{}: global zero Z0 = 0.5(ch3+ch4) at t0 of run 1 = {:.8f} m".format(test, z0))
    if test != "Test9":
        print("  NOTE masses m15/m16/m17 and mass_scale are Test9's constants, applied to the twin wall by assumption")
    files = sorted(glob.glob(os.path.join(raw_dir, test + "Run*.xlsx")), key=lambda p: int(RUN_RE.match(os.path.basename(p)).group(2)) if RUN_RE.match(os.path.basename(p)) else 0)
    n = 0
    for p in files:
        m = RUN_RE.match(os.path.basename(p))
        if not m:
            continue
        run = int(m.group(2))
        t, ch = load_raw(p)
        arr = process(t, ch, z0)
        out = os.path.join(out_dir, "{}Run{}_processed_globalzero.xlsx".format(test, run))
        write_xlsx(out, arr); n += 1
        print("  run {:2d}: {} samples, {:.2f} s, peak |U_avg| {:.3f} mm (within-run {:.3f}), peak |V| {:.2f} kN -> {}".format(
            run, len(t), t[-1], np.max(np.abs(arr[:, 1])) * 1000, np.max(np.abs(arr[:, 1] - arr[0, 1])) * 1000,
            np.max(np.abs(arr[:, 3])), os.path.basename(out)))
    print("{} file(s) written to {}".format(n, out_dir))

if __name__ == "__main__":
    main()